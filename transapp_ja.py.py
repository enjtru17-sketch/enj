#!/usr/bin/env python
# coding: utf-8

# In[1]:


import streamlit as st
from docx import Document
from deep_translator import GoogleTranslator
import io

def translate_word_file(doc, source_lang='ja', target_lang='ko'):
    """
    업로드된 Document 객체를 번역하여 반환합니다.
    """
    translator = GoogleTranslator(source=source_lang, target=target_lang)

    # 문단 번역
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    progress_bar = st.progress(0)

    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            try:
                para.text = translator.translate(para.text)
            except Exception as e:
                st.error(f"문단 번역 오류: {e}")

        # 진행률 업데이트 (문단 기준)
        if len(doc.paragraphs) > 0:
            progress_bar.progress((i + 1) / len(doc.paragraphs))

    # 테이블 번역
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip():
                        try:
                            paragraph.text = translator.translate(paragraph.text)
                        except Exception as e:
                            st.error(f"테이블 번역 오류: {e}")

    return doc

# --- Streamlit UI ---
st.set_page_config(page_title="Word Translator", page_icon="📝")
st.title("📝 워드 파일 자동 번역기")
st.write("일본어로 된 Word(.docx) 파일을 업로드하면 구글 번역기를 통해 한국어로 번역해 드립니다.")

uploaded_file = st.file_uploader("번역할 워드 파일을 선택하세요", type=["docx"])

if uploaded_file is not None:
    # 파일 로드
    doc = Document(uploaded_file)

    if st.button("번역 시작"):
        with st.spinner('번역 중입니다... 잠시만 기다려 주세요.'):
            # 번역 실행
            translated_doc = translate_word_file(doc)

            # 메모리에 파일 저장 (Streamlit 다운로드용)
            output = io.BytesIO()
            translated_doc.save(output)
            processed_data = output.getvalue()

            st.success("✅ 번역이 완료되었습니다!")

            # 다운로드 버튼 생성
            st.download_button(
                label="번역된 파일 다운로드",
                data=processed_data,
                file_name=f"translated_{uploaded_file.name}",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )



# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:




