#!/usr/bin/env python
# coding: utf-8

# In[1]:


import streamlit as st
from docx import Document
from deep_translator import GoogleTranslator
import io

def translate_word_file(doc, target_lang='ko'):
    """
    업로드된 Document 객체의 언어를 자동 감지하여 target_lang으로 번역합니다.
    """
    # source='auto'로 설정하여 입력 언어 자동 인식
    translator = GoogleTranslator(source='auto', target=target_lang)

    # 텍스트가 있는 문단/셀만 번역하여 효율성 향상
    # 문단 번역
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    progress_bar = st.progress(0)

    total_items = len(paragraphs)

    for i, para in enumerate(paragraphs):
        try:
            # 텍스트 번역
            original_text = para.text
            translated_text = translator.translate(original_text)
            para.text = translated_text
        except Exception as e:
            st.error(f"문단 번역 오류: {e}")

        # 진행률 업데이트
        if total_items > 0:
            progress_bar.progress((i + 1) / total_items)

    # 테이블 번역
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip():
                        try:
                            paragraph.text = translator.translate(paragraph.text)
                        except Exception as e:
                            st.error(f"테이블 번역 오류: {e}")

    return doc

# --- Streamlit UI ---
st.set_page_config(page_title="Word Translator", page_icon="📝")
st.title("📝 워드 파일 자동 번역기")
st.write("Word(.docx) 파일을 업로드하면 언어를 자동 감지하여 원하는 언어로 번역해 드립니다.")

# 언어 선택 사이드바
target_lang_map = {
    "한국어": "ko",
    "영어": "en",
    "일본어": "ja",
    "중국어(간체)": "zh-CN",
    "스페인어": "es",
    "러시아어": "ru"
}
target_lang_name = st.sidebar.selectbox("번역할 언어 선택", list(target_lang_map.keys()), index=0)
target_lang_code = target_lang_map[target_lang_name]

uploaded_file = st.file_uploader("번역할 워드 파일을 선택하세요", type=["docx"])

if uploaded_file is not None:
    # 파일 로드
    doc = Document(uploaded_file)

    if st.button("번역 시작"):
        with st.spinner(f"{target_lang_name} (으)로 번역 중입니다... 잠시만 기다려 주세요."):
            # 번역 실행 (자동 감지 모드)
            translated_doc = translate_word_file(doc, target_lang=target_lang_code)

            # 메모리에 파일 저장 (Streamlit 다운로드용)
            output = io.BytesIO()
            translated_doc.save(output)
            processed_data = output.getvalue()

            st.success("✅ 번역이 완료되었습니다!")

            # 다운로드 버튼 생성
            st.download_button(
                label="번역된 파일 다운로드",
                data=processed_data,
                file_name=f"translated_{target_lang_code}_{uploaded_file.name}",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )



# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:





# In[ ]:




